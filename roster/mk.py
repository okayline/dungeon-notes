import json,re
from openpyxl import Workbook
from openpyxl.styles import Font,PatternFill,Alignment
from openpyxl.utils import get_column_letter
import os,sys
OUT=os.environ['OUT_DIR']
o=json.load(open(OUT+'/rows.json'))
fam={}
for r in o['roster']:
    if r[1]=='Enemy': fam.setdefault(r[2],[]).append(r[0])
def clean(r):
    who,grp,sit,line=r
    if not line.strip(): return None
    base,_,sub=sit.partition(': ')
    if sub==base or sub=='': sit=base
    m=re.fullmatch(r'#(\d+)',sub)
    if m: sit=f"{base} (level {m.group(1)})" if base in('Noise warnings','Tired of talking') else base
    sit=re.sub(r'(Noise warnings|Tired of talking): #(\d+)',r'\1 (level \2)',sit)
    sit=re.sub(r': #\d+','',sit)
    return [who,grp,sit,line]
def rows(k): return [c for c in (clean(r) for r in o[k]) if c]
wb=Workbook(); F='Arial'
hf=Font(name=F,bold=True,color='FFFFFF'); hfill=PatternFill('solid',fgColor='1F2937')
def sheet(ws,head,data,widths,fillcol=None):
    ws.append(head)
    for r in data: ws.append(r)
    for c in ws[1]: c.font=hf; c.fill=hfill; c.alignment=Alignment(vertical='center')
    for row in ws.iter_rows(min_row=2):
        for c in row: c.font=Font(name=F); c.alignment=Alignment(wrap_text=True,vertical='top')
    for i,w in enumerate(widths,1): ws.column_dimensions[get_column_letter(i)].width=w
    ws.freeze_panes='A2'; ws.auto_filter.ref=ws.dimensions
ws=wb.active; ws.title='Roster'
sheet(ws,['Name','Type','Family','Alignment','Personality','Notes'],o['roster'],[18,22,12,12,18,50])
col={'LAW':'DBEAFE','CHAOS':'FEE2E2','NEUTRAL':'F3F4F6'}
for row in ws.iter_rows(min_row=2):
    f=col.get(row[3].value)
    if f: row[3].fill=PatternFill('solid',fgColor=f)
ws=wb.create_sheet('Party & programs')
sheet(ws,['Character','Type','Situation','Line'],rows('party'),[14,10,36,80])
ws=wb.create_sheet('Enemies (by family)')
d=[r[:3]+[r[3],', '.join(fam.get(r[0],[]))] for r in rows('enemy')]
sheet(ws,['Family','Type','Situation','Line','Enemies that use it'],d,[12,8,36,70,60])
ws=wb.create_sheet('Personalities')
sheet(ws,['Personality','Type','Situation','Line'],rows('persona'),[14,10,32,80])
ws=wb.create_sheet('Notes')
for t in ["Notes",
 "Enemies share voice lines by family (data, hardware, hybrid, folklore); the 'Enemies that use it' column lists who speaks them.",
 "{n} in a line is replaced by the speaker's name in-game.",
 "Alignment comes from family (data = LAW, hardware and hybrid = NEUTRAL, folklore = CHAOS) with a few exceptions, flagged on the Roster sheet. Programs are all LAW.",
 "Personality (PRIM, SLY, GRUFF, DREAMY, HAUGHTY) shapes how a speaker reacts to silence and to a NO; enemy families pick from the list shown on the Roster.",
 "Party members: ELF, PIXIE and CU SITH. Programs: WATCHDOG, PATCH, CACHE, SCAN, COMPILER. ATOM SLASHER is a unique enemy that never talks.",
 "Corrupted enemies get a glitched name (e.g. @ for A) and turn CHAOS; they reuse their family's lines."]:
    ws.append([t])
ws['A1'].font=Font(name=F,bold=True,size=12)
for r in ws.iter_rows(min_row=2):
    r[0].font=Font(name=F); r[0].alignment=Alignment(wrap_text=True,vertical='top')
ws.column_dimensions['A'].width=110
wb.save(os.environ['XLSX'])
print({s.title:s.max_row-1 for s in wb})
